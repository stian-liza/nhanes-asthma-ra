"""Verify the rendered report against CSV sources and its fixed code links."""
import argparse
import csv
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from pypdf import PdfReader


def rows(path):
    return list(csv.DictReader(Path(path).open(encoding='utf-8-sig')))


ap = argparse.ArgumentParser()
ap.add_argument('--output', type=Path, default=Path('outputs/client_report_20261012_v03'))
ap.add_argument('--source', type=Path, default=Path('outputs/pooled_only_20261012_v03'))
ap.add_argument('--code-ref', required=True)
a = ap.parse_args()
docx = a.output / 'NHANES_类风湿关节炎与曾患哮喘_分析报告.docx'
pdf = a.output / 'render/NHANES_类风湿关节炎与曾患哮喘_分析报告.pdf'
d = Document(docx); r = PdfReader(pdf)
assert len(r.pages) == 3 and len(d.tables) == 1
bindings = json.loads((a.output / 'report_source_bindings.json').read_text())
assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == h for p, h in bindings.items())
comparisons = 0
prevalence = {x['group']: x for x in rows(a.source / 'asthma_prevalence.csv')}
for i, key in enumerate(['No_arthritis', 'RA'], 1):
    x = prevalence[key]
    expected = [f"{int(x['n']):,}", f"{int(x['events']):,}",
                f"{100*float(x['estimate']):.2f}%\n（{100*float(x['lower']):.2f}%—{100*float(x['upper']):.2f}%）"]
    assert [cell.text for cell in d.tables[0].rows[i].cells][1:] == expected
    comparisons += 3
model_forest = ''.join(p.extract_text() for p in PdfReader(a.output / 'figures/model_forest.pdf').pages)
for x in rows(a.source / 'regression.csv'):
    for value in [f"{float(x[k]):.2f}" for k in ['OR','lower','upper']] + ['<0.001']:
        assert value in model_forest, value
        comparisons += 1
subgroups = rows(a.source / 'subgroups.csv')
for variable, filename in [('agegroup','age_forest.pdf'), ('sex','sex_forest.pdf')]:
    forest = ''.join(p.extract_text() for p in PdfReader(a.output / 'figures' / filename).pages)
    for x in [z for z in subgroups if z['subgroup_variable'] == variable]:
        for value in [f"{int(x['n']):,}"] + [f"{float(x[k]):.2f}" for k in ['OR','lower','upper']]:
            assert value in forest, value
            comparisons += 1
flow = ''.join(p.extract_text() for p in PdfReader(a.output / 'figures/sample_flow.pdf').pages)
for x in rows(a.source / 'review/sample_flow.csv')[:5]:
    assert f"{int(x['remaining']):,}" in flow
    comparisons += 1
gate = json.loads((a.source / 'review/review_gate.json').read_text())
assert f"{gate['n']:,}" in flow and f"{gate['excluded_incomplete']:,}" in flow
text = '\n'.join(p.text for p in d.paragraphs)
for x in rows(a.source / 'interaction_tests.csv'):
    assert f"{float(x['p']):.4f}" in text
    comparisons += 1
assert '未做插补、替代表型或其他敏感性分析' in text
assert '跨调查空档合并的适用性未获本研究验证' in text
assert not re.search(r'\bRA\b', text)
with ZipFile(docx) as z:
    rels = ET.fromstring(z.read('word/_rels/document.xml.rels'))
    urls = {e.attrib['Target'] for e in rels if e.attrib.get('Type', '').endswith('/hyperlink')}
pdf_urls = set()
for page in r.pages:
    for ref in page.get('/Annots', []):
        obj = ref.get_object(); action = obj.get('/A', {})
        if '/URI' in action: pdf_urls.add(str(action['/URI']))
assert urls == pdf_urls, (urls - pdf_urls, pdf_urls - urls)
assert len(urls) >= 16
for url in urls:
    if 'github.com' in url:
        assert f'/blob/{a.code_ref}/' in url
        tail = url.split(f'/blob/{a.code_ref}/')[1]
        path, _, fragment = tail.partition('#')
        assert Path(path).exists(), path
        if fragment:
            parts = [int(x) for x in re.findall(r'L(\d+)', fragment)]
            assert min(parts) >= 1 and max(parts) <= len(Path(path).read_text().splitlines())
# Guard against the initial invisible-CJK rendering failure.
for page in r.pages:
    text = page.extract_text()
    assert '类风湿关节炎' in text or '哮喘' in text
    assert '\ufffd' not in text
result = {'verdict':'pass','pages':len(r.pages),'editable_tables':len(d.tables),
          'numeric_checks':comparisons,'hyperlinks_in_docx_and_pdf':len(urls),
          'fixed_code_ref':a.code_ref,'remote_verification':'recorded separately',
          'visual_review':'required for final rendered pages'}
(a.output/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False))

"""Verify report tables, source hashes and formal-plan binding; no model execution."""
import pathlib,json,hashlib,csv
from docx import Document
ROOT=pathlib.Path(__file__).resolve().parents[2];O=ROOT/'outputs/year_extension_20261008_v01'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
b=json.loads((O/'source_bindings.json').read_text());d=Document(next(O.glob('*.docx')));checks=[]
def ck(name,condition):
 checks.append(dict(check=name,pass_check=bool(condition)))
 if not condition:raise AssertionError(name)
for file,h in b['source_tables_and_figures'].items():ck('source:'+file,sha(O/file)==h)
ck('locked_plan',b['plan_sha256']==json.loads((O/'plan_lock.json').read_text())['sha256'])
ck('table_count',len(d.tables)==len(b['tables']))
for i,(t,v) in enumerate(zip(d.tables,b['tables'])):
 ck('table:'+str(i+1),[[c.text for c in r.cells] for r in t.rows]==[[str(x) for x in v['headers']]]+[[str(x) for x in r] for r in v['rows']])
texts=[p.text for p in d.paragraphs]
for i,p in enumerate(b['prose']):ck('paragraph:'+str(i+1),p in texts)
for scope in ['prepandemic','latest','pooled']:
 ck(scope+':input',json.loads((O/scope/'independent_input_validation.json').read_text())['verdict']=='pass')
 ck(scope+':variance',json.loads((O/scope/'independent_validation.json').read_text())['all_pass'])
 ck(scope+':diagnostics',all(v=='pass' for k,v in json.loads((O/scope/'model_gate.json').read_text()).items() if k in ['G4','G5','G6']))
for scope in ['prepandemic','latest']:ck(scope+':supplement',json.loads((O/scope/'supplement/independent_validation.json').read_text())['all_pass'])
ck('pooling_variance',json.loads((O/'pooled/pooling_independent_validation.json').read_text())['all_pass'])
ck('frozen_original',sha(ROOT/'outputs/freezes/2026-09-28-asthma-ra-v01/asthma_ra_v01_frozen.zip')=='66d09e6ca2364b6967849be0f1988d96a6ad4596bf8ddff16ebb0cd23f007d31')
result=dict(verdict='pass',checks=len(checks),tables=len(d.tables),docx_sha256=sha(next(O.glob('*.docx'))),scope='source and document consistency, with independently validated new model outputs',details=checks)
(O/'report_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print('Report validation passed:',len(checks),'checks')

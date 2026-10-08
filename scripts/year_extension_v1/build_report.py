"""Create the updated editable manuscript exclusively from this run's final tables."""
import ast,csv,json,hashlib,re,pathlib,math
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
REPO=pathlib.Path(__file__).resolve().parents[2];OUT=REPO/'outputs/year_extension_20261008_v01';hashes={}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(scope,name):
 p=OUT/scope/(name+'.csv');hashes[str(p.relative_to(OUT))]=sha(p);rs=list(csv.DictReader(p.open(encoding='utf-8-sig')))
 for r in rs:
  for k,v in r.items():
   try:r[k]=float(v)
   except (ValueError,TypeError):pass
 return rs
def one(rs,**kw):return next(r for r in rs if all(r[k]==v for k,v in kw.items()))
def f(x,d=2):return f'{x:.{d}f}'
def n(x):return f'{int(x):,}'
def pct(x):return f(x*100)
def pv(x):return '<0.001' if x<.001 else f(x,3)
def ci(r,key='OR'):return f"{f(r[key])}（{f(r['lower'])}—{f(r['upper'])}）"
# Standalone formatting and verified reference metadata. No previous results are loaded.
labels={'age': '年龄/岁', 'bmi': 'BMI/（kg·m⁻²）', 'pir': 'PIR', 'agegroup': '年龄/岁', 'sex': '性别', 'race': '种族/族裔', 'education': '教育程度', 'pirgroup': 'PIR', 'bmigroup': 'BMI/（kg·m⁻²）', 'smoking': '吸烟状况', '20-39': '20～39', '40-59': '40～59', '60-79': '60～79', 'Male': '男性', 'Female': '女性', 'NH_White': '非西班牙裔白人', 'Mexican': '墨西哥裔', 'Other_Hispanic': '其他西班牙裔', 'NH_Black': '非西班牙裔黑人', 'Other_Multiracial': '其他及多种族', 'College_grad': '大学毕业及以上', 'Less_than_9th': '低于九年级', '9th_11th': '九至十一年级', 'High_school': '高中或同等学历', 'Some_college': '部分大学或副学士', 'Below_1': '＜1', 'At_least_1': '≥1', 'Underweight': '＜18.5', 'Normal': '18.5～＜25', 'Overweight': '25～＜30', 'Obese': '≥30', 'never': '从未', 'former': '既往', 'current': '当前'}
oldrefs=['SHEEN Y H, ROLFES M C, WI C I, et al. Association of asthma with rheumatoid arthritis: a population-based case-control study[J]. J Allergy Clin Immunol Pract, 2018, 6(1): 219-226. DOI: 10.1016/j.jaip.2017.06.022.', 'SPARKS J A, LIN T C, CAMARGO C A Jr, et al. Rheumatoid arthritis and risk of chronic obstructive pulmonary disease or asthma among women: a marginal structural model analysis in the Nurses’ Health Study[J]. Semin Arthritis Rheum, 2018, 47(5): 639-648. DOI: 10.1016/j.semarthrit.2017.09.005.', 'KIM J G, KANG J, LEE J H, et al. Association of rheumatoid arthritis with bronchial asthma and asthma-related comorbidities: a population-based national surveillance study[J]. Front Med (Lausanne), 2023, 10: 1006290. DOI: 10.3389/fmed.2023.1006290.', 'BALJET E, LUIJKS H, VAN DEN BEMT L, et al. Chronic comorbid conditions and asthma exacerbation occurrence in a general population sample[J]. NPJ Prim Care Respir Med, 2023, 33: 29. DOI: 10.1038/s41533-023-00350-x.', 'CENTERS FOR DISEASE CONTROL AND PREVENTION. Continuous NHANES: questionnaires, datasets, and related documentation[DB/OL]. [2026-09-29]. https://wwwn.cdc.gov/nchs/nhanes/continuousnhanes/.', 'CENTERS FOR DISEASE CONTROL AND PREVENTION. Ethics review board approval[EB/OL]. (2026-03-24)[2026-09-29]. https://www.cdc.gov/nchs/nhanes/about/erb.html.', 'CENTERS FOR DISEASE CONTROL AND PREVENTION. Adult BMI categories[EB/OL]. [2026-09-29]. https://www.cdc.gov/bmi/adult-calculator/bmi-categories.html.', 'CENTERS FOR DISEASE CONTROL AND PREVENTION. NHANES tutorial: weighting[EB/OL]. [2026-09-29]. https://wwwn.cdc.gov/nchs/nhanes/tutorials/weighting.aspx.', 'CENTERS FOR DISEASE CONTROL AND PREVENTION. NHANES tutorial: variance estimation[EB/OL]. [2026-09-29]. https://wwwn.cdc.gov/nchs/nhanes/tutorials/varianceestimation.aspx.', 'LUMLEY T. survey: analysis of complex survey samples[CP/OL]. [2026-09-29]. https://cran.r-project.org/web/packages/survey/survey.pdf.']
D=Document();s=D.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 x=D.styles[st];x.font.name='Times New Roman';x.font.color.rgb=RGBColor(0,0,0);x.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC');x.paragraph_format.space_after=Pt(4)
 for b in x.element.xpath('.//w:pBdr'):b.getparent().remove(b)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.line_spacing=1.12
D.styles['Title'].font.size=Pt(18);D.styles['Title'].font.bold=True
D.styles['Heading 1'].font.size=Pt(14);D.styles['Heading 2'].font.size=Pt(11);D.styles['Heading 2'].font.bold=True
for st in D.styles:
 if st.type==1:
  snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(snap)
for g in s._sectPr.xpath('w:docGrid'):s._sectPr.remove(g)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
D.core_properties.title='自报类风湿关节炎与曾患哮喘的关联';D.core_properties.subject='NHANES 1999—2018 横断面研究';D.core_properties.author=''
prose=[];tablebind=[]
def p(text,style=None,size=None):
 prose.append(text);q=D.add_paragraph(style=style)
 for part in re.split(r'(\[\d+(?:[-,，]\d+)*\])',text):
  rr=q.add_run(part)
  if re.fullmatch(r'\[\d+(?:[-,，]\d+)*\]',part):rr.font.superscript=True
  if size:rr.font.size=Pt(size)
 q.paragraph_format.widow_control=True
 return q
def h(text,level=1):D.add_heading(text,level)
def page(text):q=D.add_heading(text,1);q.paragraph_format.page_break_before=True

def caption(cn,en):
 q=p(cn,size=10);q.paragraph_format.keep_with_next=True
 q=p(en,size=9.5);q.paragraph_format.keep_with_next=True

def table(headers,rows,widths,size=9.5):
 t=D.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,(w,title) in enumerate(zip(widths,headers)):t.columns[j].width=Cm(w);t.rows[0].cells[j].text=title
 for row in rows:
  cc=t.add_row().cells
  for j,v in enumerate(row):cc[j].text=str(v)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tc=c._tc.get_or_add_tcPr();bor=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);show=(i==0 and edge in ['top','bottom']) or (i==len(t.rows)-1 and edge=='bottom');el.set(qn('w:val'),'single' if show else 'nil');el.set(qn('w:sz'),'10' if edge=='top' or i==len(t.rows)-1 else '5');el.set(qn('w:color'),'000000');bor.append(el)
   tc.append(bor);mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'40');el.set(qn('w:type'),'dxa');mar.append(el)
   tc.append(mar)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for rr in q.runs:rr.font.size=Pt(size);rr.bold=i==0
 tablebind.append({'headers':headers,'rows':rows});return t

D.core_properties.subject='NHANES 1999至2020年3月和2021至2023年 重复横断面研究'
prose=[];tablebind=[]
periods={'prepandemic':'1999—2020年3月','latest':'2021年8月—2023年8月','pooled':'两时段合并补充'}
P={};R={};I={};S={};B={};T={};M={}
for scope in periods:
 P[scope]=read(scope,'asthma_prevalence');R[scope]=read(scope,'regression');I[scope]=read(scope,'interaction_tests');S[scope]=read(scope,'subgroups')
 if scope!='pooled':
  B[scope]=read(scope+'/supplement','regression');T[scope]=read(scope+'/supplement','tests');M[scope]=read(scope+'/supplement','standardized_prevalence')
pre=one(R['prepandemic'],model='M3');late=one(R['latest'],model='M3');pool=read('pooled','pooling_models');p1=one(pool,model='P1',term='ra');p2=one(pool,model='P2',term='ra');i1=one(pool,model='I1',term='ra:eralatest');i2=one(pool,model='I2',term='ra:eralatest')
a=one(P['prepandemic'],group='All');b=one(P['latest'],group='All')
pa=one(P['prepandemic'],group='RA');pb=one(P['latest'],group='RA')
def picture(scope,name,width=16):
 path=OUT/scope/'figures'/name;hashes[str(path.relative_to(OUT))]=sha(path)
 q=D.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(path),width=Cm(width))

auth='［作者姓名、排序、单位、通信作者及基金信息待作者据实补齐］'
p('自报类风湿关节炎与曾患哮喘的关联','Title')
p('NHANES 1999年至2020年3月与2021年至2023年重复横断面研究','Subtitle');p(auth,size=9)
h('摘要')
p('【目的】分析自报类风湿关节炎（RA）与曾患哮喘的关联，评价其在较新调查时期的表现，并考察跨期合并估计的适用性。')
p('【方法】使用美国国家健康与营养检查调查（NHANES）1999—2020年3月及2021年8月—2023年8月资料。纳入20～79岁、疾病及七项预设协变量完整有效者。结局为自报曾被医生或其他卫生专业人员告知患哮喘；RA要求医生诊断关节炎且类型为RA；对照为明确无关节炎者。采用体检权重、分层和主要抽样单位进行子总体分析。分别建立四个复杂抽样Logistic模型；预先规定跨期合并补充、调查周期调整和时期交互检查，并进行函数形式、标准化比例及完整病例选择敏感性分析。')
p(f"【结果】疫情前及最新时期分别纳入{n(a['n'])}人和{n(b['n'])}人，其中RA者分别为{n(pa['n'])}人和{n(pb['n'])}人。完全调整模型的优势比（OR）及95%置信区间分别为{ci(pre)}和{ci(late)}。合并{n(p1['n'])}人并调整调查周期后的OR为{ci(p1,'estimate')}，加入年龄及贫困收入比非线性调整后为{ci(p2,'estimate')}。线性调整的时期交互OR比为{ci(i1,'estimate')}，P={pv(i1['p'])}；其区间仍容许一定幅度的时期差异。")
p('【结论】在规定的完整病例及无关节炎对照口径下，两个观测时期均见RA与曾患哮喘的正向横断面关联。跨期合并可提供补充性的混合人群关联估计，但不能证明时期等效、代表疫情调查空档或判断因果关系。')
p('关键词：类风湿关节炎；哮喘；横断面研究；复杂抽样；美国国家健康与营养检查调查',size=10)
page('Association Between Self Reported Rheumatoid Arthritis and Ever Diagnosed Asthma')
p('Repeated Cross Sectional Analyses of NHANES 1999 to March 2020 and August 2021 to August 2023',size=12)
p('[Author names and affiliations to be supplied]',size=9);h('Abstract',2)
p('Objective To examine the association between self-reported rheumatoid arthritis (RA) and ever-diagnosed asthma across two observed survey periods and assess supplementary pooling across periods.')
p('Methods Adults aged 20–79 years with complete disease information and seven prespecified covariates were selected from NHANES 1999–March 2020 and August 2021–August 2023. RA required self-reported doctor-diagnosed arthritis of the rheumatoid type; comparators explicitly reported no arthritis. Ever-diagnosed asthma was the outcome. Examination weights, strata and primary sampling units were incorporated into survey-domain analyses. Four fixed logistic models were fitted separately by period. Prespecified supplementary analyses addressed pooling, cycle adjustment, period interactions, nonlinear adjustment, standardized proportions and complete-case selection.')
p(f"Results The two periods included {n(a['n'])} and {n(b['n'])} participants, including {n(pa['n'])} and {n(pb['n'])} with RA, respectively. Fully adjusted odds ratios (ORs) were {f(pre['OR'])} (95% confidence interval [CI], {f(pre['lower'])}–{f(pre['upper'])}) and {f(late['OR'])} (95% CI, {f(late['lower'])}–{f(late['upper'])}). Among {n(p1['n'])} pooled participants, the cycle-adjusted OR was {f(p1['estimate'])} (95% CI, {f(p1['lower'])}–{f(p1['upper'])}); the OR with nonlinear age and poverty income ratio adjustment was {f(p2['estimate'])} (95% CI, {f(p2['lower'])}–{f(p2['upper'])}). The latest-to-prepandemic OR ratio in the linear-adjustment interaction model was {f(i1['estimate'])} (95% CI, {f(i1['lower'])}–{f(i1['upper'])}; P={pv(i1['p'])}), leaving uncertainty about differences between periods.")
p('Conclusion Positive cross-sectional associations were observed in both periods under the specified complete-case and no-arthritis comparator definitions. Pooling provides a supplementary association for a mixture of observed periods, without establishing equivalence, representing the unobserved pandemic interval, or identifying causality.')
p('Keywords: rheumatoid arthritis; asthma; cross-sectional study; complex survey; NHANES',size=10)
page('引言')
p('RA与哮喘的共存可能涉及共同危险因素、诊断机会和治疗等多种解释。既往病历病例对照及全国调查报告过两病关联[1,3]；女性前瞻性研究在控制时间变化因素后，未发现RA与后续新发哮喘的明确关联[2]。曾患哮喘、当前哮喘发作及新发哮喘并非相同结局[4]，研究设计和疾病定义差异应纳入解释。')
p('本研究延续已确认的疾病与对照定义，将疫情前资料扩展至2020年3月，并增加最新公开的2021—2023年资料。NHANES疫情期间暂停调查，CDC一般不建议将最新周期直接与早期周期合并[5]。因此，本研究保留分时期估计，同时将跨期合并作为预先规定的补充，以区分总体关联、时期差异及测量限制。')
h('1 资料与方法')
h('1.1 数据来源与人群',2)
p('本研究为公开资料的重复横断面二次分析。使用DEMO、MCQ、SMQ、BMX四模块，全新获取44个官方数据文件，按受访者编号一对一连接。疫情前由1999—2016年九个两年周期和2017—2020年3月特殊发布块组成；后者已包含2017—2018年资料，因此未重复纳入旧2017—2018年文件。最新时期单独使用2021年8月—2023年8月资料[5-7]。不同调查时期不是同一批受访者的随访。')
p('纳入年龄20～79岁、体检权重和设计变量有效、两病状态明确且七项协变量完整者。其他关节炎类型、类型未知、拒答、不知道、疾病或协变量实际缺失分别记录并排除。无关节炎者类型题的合法跳题和从未吸烟者后续题的合法跳题不视为缺失。最终筛选顺序及各步人数见图1、图2及完整筛选表。')
p('原调查具有NHANES伦理审查批准和参与者知情同意[8]。本研究使用去标识化公开资料。［本单位公开数据二次研究的伦理审查或豁免认定、编号待据实补齐］。')
h('1.2 疾病及协变量',2)
p('曾患哮喘由MCQ010肯定回答界定，不要求当前仍患哮喘。RA要求MCQ160A肯定且关节炎类型为RA；明确否定关节炎总题者作对照。1999—2008年类型字段为MCQ190，2009—2010年为MCQ191，两者RA编码为1；此后采用MCQ195，RA编码为2。未知和其他关节炎类型不编码为对照。逐发布块核对题干、代码本、适用年龄、频数及问卷跳题。')
p('协变量为年龄、性别、种族/族裔、教育、贫困收入比（PIR）、体质指数（BMI）及吸烟。种族采用各期共同的RIDRETH1五类，教育为成人五级。BMI采用官方实测值；PIR为有效0～5值，5保留顶码含义。吸烟按终生100支及当前吸烟题分为从未、既往、当前。最新问卷使用SMQ.022及产品提示卡，发布字段仍为SMQ020；据公开题干保留同一100支分类，提示方式差异记为测量限制。')
p('分类比较年龄分20～39、40～59、60～79岁；成人BMI分＜18.5、18.5～＜25、25～＜30、≥30 kg/m²[9]；PIR分＜1和≥1。主回归三者均连续线性调整。采用完整病例；不实施多重插补或因极端值本身而截尾。')
h('1.3 复杂抽样与权重',2)
p('因模型包含实测BMI，采用体检权重。疫情前21.2年分析：1999—2002年使用WTMEC4YR乘4/21.2，2003—2016年使用WTMEC2YR乘2/21.2，2017—2020年3月使用WTMECPRP乘3.2/21.2。最新周期单独采用WTMEC2YR。上述依据CDC四年权重及特殊3.2年发布块规则构建[5,6]。')
p('跨期补充定义为按已观测发布块时长加权的混合人群，总时长23.2年；四年块乘4/23.2，普通两年块及最新周期乘2/23.2，特殊块乘3.2/23.2。该定义不补回2020年4月—2021年7月空档，也不是CDC对跨疫情合并的特别认可。不得把混合样本估计写作2023年当前水平。')
p('纳入官方分层SDMVSTRA和层内抽样单位SDMVPSU，在有效MEC宽样本建立设计，再取目标完整病例及亚组子总体；采用Taylor线性化方差，保留父设计抽样信息，不无依据拼接早期周期分层号[10]。R及survey的实际版本见复现附件。单系数以子总体设计自由度进行t检验及区间推断，联合检验采用设计Wald F方法[11]。')
h('1.4 模型与预设补充',2)
p('每个分析集合四模型使用同一完整病例：M0未调整；M1增加年龄、性别、种族；M2再增加教育和PIR；M3再增加BMI及吸烟。主结果为M3。分类比较采用Rao–Scott校正F检验。年龄和性别亚组维持预设分组，组内调整其他协变量；交互检验来自总体模型，不能比较各组显著与否。检验为双侧，辅助分析未作多重性校正，解释结合效应大小、区间及预设地位。')
p('补充B0等同M3；B1以三节点限制性立方样条调整年龄与PIR；B2再增加BMI样条；B3在B1增加发布块因子，仅适用于多周期的疫情前样本；B4以原体检权重除以完整病例纳入概率拟合B1，联合方差计入选择概率估计不确定性。选择模型使用两病、年龄及其非线性、性别、种族与可变的调查周期，不验证不可观察缺失机制。')
kn=read('prepandemic/supplement','knots')
p('样条节点在本轮拟合前按疫情前完整病例自变量第10、50、90百分位锁定，最新周期和合并分析共用：'+ '；'.join(labels[r['variable']]+' '+', '.join(f(r[k]) for k in ['q10','q50','q90']) for r in kn)+'。标准化比例以各自时间集合完整病例的加权协变量分布为参考，方差同时包含参数及参考分布抽样波动；不同集合的标准化比例不能直接作共同标准人口下的趋势比较。')
p('跨期合并补充包括M0—M3、P1（M3增加发布块因子）、P2（P1增加年龄和PIR非线性项）。在P1、P2分别加入RA与时期乘积项；时期主效应已由发布块因子覆盖。交互系数指数变换为最新时期与疫情前OR之比。交互模型假定其他协变量效应共同，另保留两时期独立模型。不以交互不显著认定等效，不依据结果选择主分析。')
p('检查唯一编号、官方频数、筛选闭合、单元格、收敛、设计矩阵秩、完全或准完全分离、共线性及精度。独立重建原始字段、资格与权重，独立PSU汇总计算设计方差，并以等价样条基、原生预测及显式Wald公式核验。一次补充方差检查发现默认收敛精度不足后，将求解容差收紧至1×10⁻¹²并重跑，未放宽校验阈值。')
h('2 结果')
h('2.1 人群及两组曾患哮喘比例',2)
for scope in ['prepandemic','latest']:
 q=one(P[scope],group='RA');c=one(P[scope],group='No_arthritis');fl=read(scope,'sample_flow')
 p(f"{periods[scope]}最终纳入{n(q['n']+c['n'])}人，其中RA组{n(q['n'])}人、无关节炎组{n(c['n'])}人；曾患哮喘者分别为{n(q['events'])}人和{n(c['events'])}人，加权比例分别为{pct(q['estimate'])}%和{pct(c['estimate'])}%（表1、表2）。")
h('2.2 主要回归及亚组',2)
p(f"疫情前M3的OR为{ci(pre)}，最新时期为{ci(late)}（表3）。两者均为正向关联，最新时期区间较宽。两时间集合实际设计自由度分别为163和15；其差异反映抽样设计支持，不能以总人数代替。")
for scope in ['prepandemic','latest']:
 ai=one(I[scope],variable='agegroup');si=one(I[scope],variable='sex');nai=one(T[scope],test='B1_interaction_agegroup')
 p(f"{periods[scope]}年龄和性别交互P值分别为{pv(ai['p'])}和{pv(si['p'])}；年龄及PIR非线性调整后年龄交互P={pv(nai['p'])}（表4及补充表）。年龄或性别差异均不能仅据单个亚组显著性判断。")
h('2.3 合并补充及时期差异',2)
p(f"两时段合并共{n(p1['n'])}人。P1的OR为{ci(p1,'estimate')}，P2为{ci(p2,'estimate')}（表5）。线性及非线性调整下，最新时期相对于疫情前的OR比分别为{ci(i1,'estimate')}及{ci(i2,'estimate')}，交互P值分别为{pv(i1['p'])}和{pv(i2['p'])}。这些区间未排除一定幅度的时期差异；合并估计保留补充地位。")
h('2.4 函数形式与完整病例选择',2)
for scope in ['prepandemic','latest']:
 b1=one(B[scope],model='B1');b4=one(B[scope],model='B4');rd=one(M[scope],model='B1',estimand='difference_RA_minus_control');sw=read(scope+'/supplement','selection_weight_diagnostics')[0]
 p(f"{periods[scope]}B1的OR为{ci(b1)}，B4为{ci(b4)}；B1标准化比例差为{pct(rd['estimate'])}个百分点（95%CI {pct(rd['lower'])}—{pct(rd['upper'])}）。预测完整病例概率范围为{f(sw['min_probability'],3)}—{f(sw['max_probability'],3)}；选择加权结果依赖明确假设，不证明选择偏倚已消除。")
p('模型数值资格及独立核验通过；稀疏类别和较宽区间仍保留。连续变量与各亚组的详细频数、缺失模式、共同支持和模型诊断均在汇总附件中提供。部分老年亚组样条基相关较高，等价正交基验证用于排除数值表示造成的差异，不用于删变量。')
h('3 讨论')
p('更新后两个观测时期均支持自报RA与曾患哮喘的正向关联；主要估计对预设函数形式和选择权重处理相近。这为两病共存提供新的时期资料。它与病例对照及横断面文献方向相容[1,3]，但不能替代前瞻性证据；既往RA与新发哮喘关联不明确的纵向结果仍需纳入考虑[2]。')
p('合并问题需要分别评价技术资格与科学解释。文件、疾病口径、权重及模型检查通过，说明本次预设混合样本参数可以计算。分时期估计方向一致、周期调整前后相近，有助于理解这一补充参数。但时期OR比的区间仍较宽，最新周期抽样自由度较低，不能用交互P值大于0.05证明关联稳定或两时期等效；疫情空档的分布始终未被观测[5,7]。')
p('分析直接支持指定样本中的统计关联（direct evidence）。自报医生诊断是临床疾病的替代测量（proxy evidence），可能受到关节炎类型辨认、回忆及诊断机会影响。共同炎症、药物作用或免疫机制仅为可能解释（speculation），本研究未测定其作用。纵向研究中的不明确结果、时期区间允许的差异及稀疏亚组属于应保留的反向或不确定证据（counterevidence/uncertainty）。')
p('首先，横断面资料无法确定先后顺序，OR不能表述为新发风险比或因果效应。其次，对照仅为明确无关节炎者，排除了其他类型及未知者，不能推广为RA对所有非RA的比较。两种疾病均未通过病历、血清学或肺功能逐例确诊，未测疾病活动度、治疗、就医频率等可造成残余混杂。')
p('再次，完整病例可能具有选择性；选择概率模型只能评价给定可观察信息的一种解释。最新周期访问方式、抽样设计及应答情况改变，吸烟题提示方式与RA帮助说明存在措辞差别，即使字段编码协调，也不能排除测量差异。不同标准人口下的标准化比例差不可直接解释为时期变化。多项辅助检验未经多重性校正，亚组结果应审慎解释。')
p('最后，合并补充代表已观测时段的特定混合人群，不能反映缺失疫情时段或2026年当前水平。结果来自美国成人，不能直接代表中国或宝安区；数据未包含中医证候或中医药干预，不能推断证型或治疗效果。后续应优先加强疾病确认及时间顺序研究，不以添加复杂方法替代这些证据。')
h('4 结论')
p('在规定的完整病例及明确无关节炎对照口径下，疫情前与2021—2023年均观察到自报RA与曾患哮喘的正向横断面关联。预设合并分析可作为补充，但分时期结果和不确定性应同时报告；目前不能证明时期等效、因果关系或任何治疗获益。')
p('作者贡献、利益冲突、基金、伦理认定及代码公开存档：［待作者据实补齐并确认］。',size=9)
h('参考文献')
refs=list(dict.fromkeys(oldrefs[:4]))
# The original list is checked by title so repeated entries cannot shift numbering.
refs=[next(r for r in oldrefs if r.startswith(k)) for k in ['SHEEN','SPARKS','KIM','BALJET']]
refs += [
'CENTERS FOR DISEASE CONTROL AND PREVENTION. NHANES tutorial: weighting[EB/OL]. [2026-10-08]. https://wwwn.cdc.gov/nchs/nhanes/tutorials/weighting.aspx.',
'CENTERS FOR DISEASE CONTROL AND PREVENTION. NHANES 2017–March 2020 prepandemic data[DB/OL]. [2026-10-08]. https://wwwn.cdc.gov/nchs/nhanes/continuousnhanes/default.aspx?Cycle=2017-2020.',
'CENTERS FOR DISEASE CONTROL AND PREVENTION. Brief overview and analytic guidelines for NHANES August 2021–August 2023[EB/OL]. [2026-10-08]. https://wwwn.cdc.gov/nchs/nhanes/continuousnhanes/overviewbrief.aspx?Cycle=2021-2023.',
next(r for r in oldrefs if 'Ethics review' in r),next(r for r in oldrefs if 'Adult BMI' in r),next(r for r in oldrefs if 'variance estimation' in r),next(r for r in oldrefs if r.startswith('LUMLEY'))]
for i,r in enumerate(refs,1):p(f'[{i}] '+r,size=9)

for scope,fig_no in [('prepandemic',1),('latest',2)]:
 page(f'图{fig_no} {periods[scope]}样本筛选流程');p(f'Figure {fig_no} Participant selection in '+('1999–March 2020' if scope=='prepandemic' else 'August 2021–August 2023'),size=9);picture(scope,'sample_flow.png',15)
 p('注：顺序排除人数避免重复计数；疾病未知、其他关节炎及各项协变量缺失的细分人数见同名样本筛选CSV。',size=9)
for scope in ['prepandemic','latest']:
 page('表1 '+periods[scope]+'人群特征');p('Table 1 Participant characteristics by arthritis status',size=9)
 chars=read(scope,'characteristics');rs=[]
 for r in [x for x in chars if x['group']=='All']:
  vals=[]
  for g in ['All','No_arthritis','RA']:
   x=one(chars,group=g,variable=r['variable'],level=r['level']);vals.append(f"{f(x['estimate'])}（{f(x['se'])}）" if r['level']=='continuous' else f"{n(x['n'])}（{pct(x['estimate'])}）")
  label=labels[r['variable']] if r['level']=='continuous' else labels[r['variable']]+' '+labels[r['level']];rs.append([label]+vals)
 table(['特征','总体','无关节炎','RA'],rs,[6.1,3.6,3.6,3.7],9)
 p('注：连续量为加权均值（标准误），分类量为实际人数（加权百分比）；表内RA占比不代表全部美国成人RA患病率。两时期分别列示。',size=9)
page('表2 各组曾患哮喘加权比例');p('Table 2 Weighted proportions of ever-diagnosed asthma',size=9)
table(['时期','组别','病例数/人数','加权比例及95%CI/%'],[[periods[s],{'All':'总体','RA':'RA','No_arthritis':'无关节炎'}[r['group']],n(r['events'])+'/'+n(r['n']),pct(r['estimate'])+'（'+pct(r['lower'])+'—'+pct(r['upper'])+'）'] for s in ['prepandemic','latest'] for r in P[s]],[4.4,2.2,4,6.4],9.5)
p('注：采用复杂抽样logit区间，实际人数与加权比例分别报告。',size=9)
h('表3 预设四模型回归结果');p('Table 3 Prespecified survey-weighted logistic models',size=9)
table(['时期','模型','人数','OR（95%CI）','P值'],[[periods[s],r['model'],n(r['n']),ci(r),pv(r['p'])] for s in periods for r in R[s]],[4.3,1.3,2.1,6.2,3.1],9)
p('注：同一时期内四模型样本一致；M0至M3调整顺序见方法。合并四模型均为补充，周期调整见表5。',size=9)
page('表4 年龄与性别亚组');p('Table 4 Prespecified age and sex subgroups',size=9)
table(['时期','亚组','人数','OR（95%CI）','P值'],[[periods[s],labels[r['subgroup']],n(r['n']),ci(r),pv(r['p'])] for s in ['prepandemic','latest'] for r in S[s]],[4.3,2.1,2.1,5.4,3.1],9)
p('注：年龄组内继续调整连续年龄；性别组内移除恒定性别。交互检验见正文。',size=9)
for scope,j in [('prepandemic',3),('latest',4)]:
 page(f'图{j} {periods[scope]}亚组森林图');p(f'Figure {j} Age and sex subgroup estimates',size=9);picture(scope,'subgroup_forest.png',17)
 p('注：点及横线为OR与95%CI，横轴为对数刻度；虚线为OR=1。组内显著性不能代替交互检验。',size=9)
page('表5 跨期合并与时期交互');p('Table 5 Supplementary pooling and period interactions',size=9)
rs=[]
for r in pool:
 if r['model'] in ['P1','P2'] or r['term']=='ra:eralatest':rs.append([r['model'],'总体OR' if r['term']=='ra' else '最新/疫情前OR比',n(r['n']),ci(r,'estimate'),pv(r['p'])])
table(['模型','参数','人数','估计（95%CI）','P值'],rs,[1.4,4.2,2,6.4,3],9.5)
p('注：P1为M3加周期；P2再加年龄/PIR非线性。I1/I2在P1/P2中加入RA与时期交互。OR比为1表示模型中的关联相同；其置信区间用于评估不确定性。此表不检验缺失疫情时段，也不构成等效性证明。',size=9)
h('表6 分时期补充模型');p('Table 6 Period-specific supplementary models',size=9)
table(['时期','模型','OR（95%CI）','P值'],[[periods[s],r['model'],ci(r),pv(r['p'])] for s in ['prepandemic','latest'] for r in B[s]],[4.6,2,7.2,3.2],9.5)
p('注：B0—B4含义见方法；最新时期只有一个发布块，B3不适用。B4区间计入选择概率估计不确定性。',size=9)
page('补充表S1 标准化比例与绝对差');p('Supplementary Table S1 Standardized proportions and absolute differences',size=9)
for s in ['prepandemic','latest']:
 h(periods[s],2)
 table(['模型','比较量','估计及95%CI/%或百分点'],[[r['model'],{'No_arthritis':'无关节炎','RA':'RA','difference_RA_minus_control':'RA减无关节炎'}[r['estimand']],pct(r['estimate'])+'（'+pct(r['lower'])+'—'+pct(r['upper'])+'）'] for r in M[s]],[1.6,5.6,9.8],9)
p('注：每个时期以其自身加权协变量分布为参考；绝对差为模型标准化的关联描述，不是干预效应。',size=9)
for s in ['prepandemic','latest']:
 page('补充表S2 '+periods[s]+'协变量与两病分布');p('Supplementary Table S2 Covariates and disease distributions',size=9)
 ca=read(s,'covariates_asthma');cr=read(s,'covariates_ra');rows=[]
 for r in ca:
  x=one(cr,variable=r['variable'],level=r['level']);rows.append([labels[r['variable']]+' '+labels[r['level']],n(r['n']),n(r['events'])+'（'+pct(r['estimate'])+'）',n(x['events'])+'（'+pct(x['estimate'])+'）'])
 table(['类别','人数','哮喘病例（加权%）','RA病例（加权%）'],rows,[6.5,2.1,4.2,4.2],9)
 p('注：比例分母为该类别的研究样本；RA比例仅适用于RA与无关节炎的选定比较。',size=9)
page('补充表S3 分类比较的设计校正检验');p('Supplementary Table S3 Design-adjusted categorical tests',size=9)
for s in ['prepandemic','latest']:
 if s=='latest':page('补充表S3（续） '+periods[s])
 else:h(periods[s],2)
 ts=read(s,'covariate_tests')
 table(['结局','变量','F值','分子/分母自由度','P值'],[[{'asthma':'哮喘','ra':'RA'}[r['outcome']],labels[r['variable']],f(r['F'],3),f(r['df1'],2)+'/'+f(r['df2'],2),pv(r['p'])] for r in ts],[2,3.7,2.7,5.6,3],9)
p('注：Rao–Scott F校正，未作多重性校正；各项为总体分类比较，不能解释为独立危险因素。',size=9)
page('补充表S4 非线性亚组与检验');p('Supplementary Table S4 Nonlinear-adjustment subgroups and tests',size=9)
for s in ['prepandemic','latest']:
 h(periods[s],2);rs=read(s+'/supplement','all_estimates');sub=[r for r in rs if r['model'].startswith(('agegroup_','sex_','period_'))]
 table(['子总体','人数','OR（95%CI）','P值'],[[r['model'].replace('agegroup_','').replace('sex_','').replace('period_','').replace('Male','男性').replace('Female','女性').replace('2009-March2020','2009—2020年3月'),n(r['n']),ci(r),pv(r['p'])] for r in sub],[5,2,7,3],9)
 for r in T[s]:
  testlabels={'B1_joint_nonlinearity':'年龄及PIR联合非线性','B1_age_nonlinearity':'年龄非线性','B1_pir_nonlinearity':'PIR非线性','B2_bmi_nonlinearity':'BMI额外非线性','B1_interaction_agegroup':'年龄组交互','B1_interaction_sex':'性别交互','B3_interaction_period':'疫情前内部时期交互'}
  ps=('P'+pv(r['p'])) if r['p']<.001 else ('P='+pv(r['p']))
  p(f"{testlabels.get(r['test'],r['test'])}：F={f(r['F'],3)}，自由度{n(r['df1'])}/{n(r['df2'])}，{ps}。",size=9)
page('补充表S5 完整病例选择与缺失')
p('Supplementary Table S5 Complete-case selection and missingness',size=9)
for s in ['prepandemic','latest']:
 h(periods[s],2);mi=read(s,'missingness');sw=read(s+'/supplement','selection_weight_diagnostics')[0];sp=read(s+'/supplement','support_overlap')[0]
 table(['协变量','疾病资格分母','缺失人数'],[[labels.get(r['variable'],r['variable']),n(r['denominator']),n(r['missing'])] for r in mi],[7,5,5],9)
 p(f"选择概率{f(sw['min_probability'],3)}—{f(sw['max_probability'],3)}，最大选择权重倍数{f(sw['max_multiplier'],3)}；协变量预测RA概率的共同支持范围外{n(sp['unweighted_outside_n'])}人，加权占比{pct(sp['weighted_outside_fraction'])}%。",size=9)
 p('各缺失人数允许重叠；顺序排除见流程表。纳入/排除者可观察特征及完整缺失组合见附件selection_characteristics.csv和missing_patterns.csv，均保留各自分母。',size=9)
name='哮喘与类风湿关节炎_更新至2023年_研究稿_v03.docx';D.save(OUT/name)
(OUT/'manuscript_text.txt').write_text('\n\n'.join(prose))
(OUT/'source_bindings.json').write_text(json.dumps({'source_tables_and_figures':hashes,'tables':tablebind,'prose':prose,'new_models_run':True,'plan_sha256':sha(REPO/'docs/research_plans/2026-10-08-asthma-ra-year-extension-experiment-plan.md'),'references':refs},ensure_ascii=False,indent=2))
print('Created',OUT/name,'tables',len(tablebind))
